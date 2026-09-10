import json
from openpyxl import load_workbook
from openai import OpenAI

client = OpenAI(api_key="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") # Use the same API setup as your project

def generate_audit_summary(input_file, requirement):

    wb = load_workbook(input_file, data_only=False)

    workbook_info = {
        "Workbook Name": input_file.name,
        "Total Sheets": len(wb.sheetnames),
        "Sheet Names": wb.sheetnames
    }

    prompt = f"""
You are an EUC Audit Expert.

Workbook Information:
{json.dumps(workbook_info, indent=2)}

Business Requirement:
{requirement}

Generate an audit summary.

Return ONLY valid JSON.

{{
  "audit_summary":"",
  "business_requirement":"",
  "workbook_name":"",
  "total_sheets":"",
  "risk_level":"",
  "change_type":"",
  "compliance_status":"",
  "recommendation":""
}}
"""

    response = client.chat.completions.create(
        model="gpt-4.1-mini", # Use the same model as your project
        messages=[
            {"role":"system","content":"You are an EUC Audit Expert."},
            {"role":"user","content":prompt}
        ],
        temperature=0
    )

    return json.loads(response.choices[0].message.content)