import json
from openpyxl import load_workbook
from openai import OpenAI

client = OpenAI(api_key="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") # Replace with your API key setup

def impact_analysis(input_file, requirement):

    # Read workbook directly from the uploaded file
    wb = load_workbook(input_file, data_only=False)

    workbook_info = {
        "sheet_names": wb.sheetnames,
        "total_sheets": len(wb.sheetnames)
    }

    prompt = f"""
You are an Excel Impact Analysis Expert.

Workbook Information:
{json.dumps(workbook_info, indent=2)}

Business Requirement:
{requirement}

Return ONLY valid JSON in this format:

{{
  "summary": "",
  "affected_sheets": [],
  "affected_columns": [],
  "affected_formulas": [],
  "dependencies": [],
  "risk_level": "Low",
  "recommendation": ""
}}
"""

    response = client.chat.completions.create(
        model="gpt-4.1-mini", # Use the same model as your Requirement Analyzer
        messages=[
            {
                "role": "system",
                "content": "You are an Excel Impact Analysis Expert."
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        temperature=0
    )

    result = response.choices[0].message.content

    return json.loads(result)
