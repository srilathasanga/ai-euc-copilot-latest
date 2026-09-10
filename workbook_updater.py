from openpyxl import load_workbook


def update_workbook(input_file, output_file, analysis):
    """
    Generic Workbook Updater
    """

    wb = load_workbook(input_file)

    operation = analysis.get("operation", "").lower()
    target = analysis.get("target", {})
    params = analysis.get("parameters", {})

    sheet_name = target.get("sheet")

    # If sheet is not specified, use first sheet
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb[wb.sheetnames[0]]

    # ---------------------------------------------------------
    # ADD SHEET
    # ---------------------------------------------------------
    if operation == "add_sheet":

        sheet_name = params.get("new_sheet", "NewSheet")
        columns = params.get("columns", [])

        # does the wb already have a sheet check?
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(title=sheet_name)
        else:
            ws = wb[sheet_name]

        # Support both list and comma-separated string
        if isinstance(columns, str):
            columns = [c.strip() for c in columns.split(",") if c.strip()]

        # Add column headers
        for col_index, column_name in enumerate(columns, start=1):
            ws.cell(row=1, column=col_index).value = column_name

    # ---------------------------------------------------------
    # DELETE SHEET
    # ---------------------------------------------------------
    elif operation == "delete_sheet":

        if ws.title in wb.sheetnames:
            wb.remove(ws)

    # ---------------------------------------------------------
    # RENAME SHEET
    # ---------------------------------------------------------
    elif operation == "rename_sheet":

        new_name = params.get("new_name")

        if new_name:
            ws.title = new_name

    # ---------------------------------------------------------
    # ADD COLUMN
    # ---------------------------------------------------------
    elif operation == "add_column":

        target = analysis.get("target", {})

        sheet_name = target.get("sheet")
        new_column = target.get("column_name")
        after_column = target.get("after_column")

        ws = wb[sheet_name]

        # Default: add at end
        insert_at = ws.max_column + 1

        # Find the column after which to insert
        if after_column:
            for cell in ws[1]:
                if str(cell.value).strip().lower() == after_column.strip().lower():
                    insert_at = cell.column + 1
                    break

        # Insert the new column
        ws.insert_cols(insert_at)

        # Set only the header
        ws.cell(row=1, column=insert_at).value = new_column

        # Leave all remaining cells blank
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=insert_at).value = ""

        print(f"Added column '{new_column}' after '{after_column}' in sheet '{sheet_name}'.")



    # ---------------------------------------------------------
    # DELETE COLUMN
    # ---------------------------------------------------------
    elif operation == "delete_column":

        column_name = target.get("column")

        for cell in ws[1]:

            if str(cell.value).strip().lower() == column_name.lower():

                ws.delete_cols(cell.column)
                break

    # ---------------------------------------------------------
    # INSERT ROW
    # ---------------------------------------------------------
    elif operation == "insert_row":

        row_no = int(target.get("row", ws.max_row + 1))

        ws.insert_rows(row_no)

    # ---------------------------------------------------------
    # DELETE ROW
    # ---------------------------------------------------------
    elif operation == "delete_row":

        row_no = int(target.get("row"))

        ws.delete_rows(row_no)

    # ---------------------------------------------------------
    # UPDATE CELL
    # ---------------------------------------------------------
    elif operation == "update_cell":

        cell = target.get("cell")

        new_value = params.get("new_value")

        if cell:
            ws[cell] = new_value

    # ---------------------------------------------------------
    # UPDATE FORMULA
    # ---------------------------------------------------------
    elif operation.lower() == "update_formula":


        sheet_name = analysis.get("sheet")
        formula_template = params.get("formula")
        #debug steps
        print("Operation", operation),
        print("Sheet Name", sheet_name),
        print("Formula Template", formula_template),


        if not sheet_name:
            raise ValueError("Sheet name is missing.")

        if not formula_template:
            raise ValueError("Formula template is missing.")

        ws = wb[sheet_name]

        for row in ws.iter_rows():
            for cell in row:

                # Update only formula cells
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    # Remove '=' from existing formula
                    old_formula = cell.value[1:]

                    # Replace placeholder with the existing formula
                    new_formula = formula_template.replace(
                        "existing_formula",
                        f"({old_formula})"
                    )

                    # Write updated formula
                    cell.value = "=" + new_formula

    # ---------------------------------------------------------
    # COPY COLUMN
    # ---------------------------------------------------------
    elif operation == "copy_column":

        source = target.get("column")

        new_column = params.get("new_column")

        src_col = None

        for cell in ws[1]:

            if str(cell.value).strip().lower() == source.lower():
                src_col = cell.column
                break

        if src_col:

            insert_col = ws.max_column + 1

            ws.insert_cols(insert_col)

            ws.cell(1, insert_col).value = new_column

            for r in range(2, ws.max_row + 1):

                ws.cell(r, insert_col).value = ws.cell(r, src_col).value

    # ---------------------------------------------------------
    # RENAME COLUMN
    # ---------------------------------------------------------
    elif operation == "rename_column":

        old_name = target.get("column")

        new_name = params.get("new_name")

        for cell in ws[1]:

            if str(cell.value).strip().lower() == old_name.lower():

                cell.value = new_name
                break

    else:

        print("Unsupported operation:", operation)

    wb.save(output_file)

    return output_file
