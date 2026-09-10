import streamlit as st
import pandas as pd
from workbook_reader import read_workbook, workbook_information,get_formula_information
from requirement_analyzer import analyze_requirement
from impact_analyzer import impact_analysis
from workbook_updater import update_workbook
from change_request_generator import generate_change_request
from docx import Document
from io import BytesIO

from audit_summary import generate_audit_summary
from risk_score import calculate_risk_score

st.title("AI EUC Change Management Copilot")

# -----------------------------
# Upload Workbook
# -----------------------------
uploaded_file = st.file_uploader(
    "Upload Production Workbook",
    type=["xlsx", "xlsm", "csv"]
)

if uploaded_file is not None:

    # Read workbook
    workbook = read_workbook(uploaded_file)

    # Workbook Information
    info = workbook_information(workbook, uploaded_file)
    # to add below line existing formula info will show in app
    formula_info = get_formula_information(workbook)
    # st.write(formula_info)-->to debug structure in side and display it in streamlit app
#--------display workbook info------------------
    st.subheader("Workbook Information")

    st.write("**Workbook Name:**", info["Workbook Name"])
    st.write("**Total Sheets:**", info["Total Sheets"])

    st.write("**Sheet Names:**")
    for sheet in info["Sheet Names"]:
        st.write("•", sheet)
    # -----------------------------
    # Display formula info
    # -----------------------------
    st.subheader("Formula Information")

    formula_info = get_formula_information(workbook)

    if formula_info:

        df = pd.DataFrame(formula_info)

        csv = df.to_csv(index=False).encode("utf-8")

        st.success("Formula report generated successfully.")

        st.download_button(
            label="📥 Download Formula Information Report",
            data=csv,
            file_name="Formula_Information_Report.csv",
            mime="text/csv"
        )

    else:
        st.info("No formulas found in the workbook.")

    # -----------------------------
    # Requirement Analyzer
    # -----------------------------
    st.subheader("Requirement Analyzer")

    requirement = st.text_area(
        "Enter Business Requirement"
    )

    if st.button("Analyze Requirement"):

        result = analyze_requirement(requirement)
        # st.write(result)  debug structure
        # st.success("Requirement analyzed successfully!")

        # Save result in session state
        st.session_state.result = result

    # -----------------------------
    # Display Analysis - debug imp
    # -----------------------------
    if "result" in st.session_state:

        result = st.session_state.result

        st.success("Requirement analyzed successfully!")
    #
    #     st.subheader("Requirement Analysis")
    #
    #     st.write("**Operation:**", result.get("operation", ""))
    #
    #     st.write("**Sheet:**", result.get("sheet", ""))
    #
    #     st.write("**Target:**")
    #     st.json(result.get("target", {}))
    #
    #     st.write("**Parameters:**")
    #     st.json(result.get("parameters", {}))
    #
    #     st.write("**Summary:**", result.get("summary", ""))
    #
    #     st.write("requirement result:",result)
    #     st.session_state.result = result

        # ===================================================
        # IMPACT ANALYSIS
        # ===================================================

        # st.subheader("📊 Impact Analysis")
        #
        # if st.button("Run Impact Analysis"):
        #
        #     if uploaded_file is None:
        #         st.error("Please upload a Production Workbook first.")
        #
        #     elif result is None:
        #         st.error("Please analyze the Business Requirement first.")
        #
        #     else:
        #
        #         # Show Requirement Analyzer Output
        #         # st.write("### Requirement Analysis Output")
        #         # st.json(result)
        #
        #         # Run Impact Analysis
        #         impact = impact_analysis(uploaded_file, result)
        #         #calculate AI Risk score
        #         risk = calculate_risk_score(result, impact)
        #
        #         st.success("Impact Analysis Completed Successfully")
        #
        #         st.write("## Impact Analysis Result")
        #
        #         st.write("**Operation:**", impact["operation"])
        #         st.write("**Target:**", impact["target"])
        #         st.write("**Affected Sheet:**", impact["affected sheet"])
        #         st.write("**Total Rows:**", impact["total rows"])
        #         st.write("**Total Columns:**", impact["total columns"])
        #         st.write("**Impact Level:**", impact["impact level"])
        #         st.write("**Recommendation:**", impact["recommendation"])
        #         # ==========================
        #         # AI Risk Score
        #         # ==========================
        #
        #         # =========================================================
        #         # AI Risk Score
        #         # =========================================================
        #
        #         risk = calculate_risk_score(result, impact)
        #
        #         st.markdown("---")
        #         st.markdown("# 🤖 AI Risk Score")
        #
        #         score = risk.get("risk_score", 0)
        #         level = risk.get("risk_level", "Low")
        #         reasons = risk.get("reason", [])
        #
        #         # Display Risk Score
        #         st.markdown(
        #             f"""
        #             <h1 style='text-align:center;
        #                        color:#1f77b4;
        #                        font-size:55px;'>
        #                 {score}/100
        #             </h1>
        #             """,
        #             unsafe_allow_html=True
        #         )
        #
        #         # Display Risk Level
        #         if level == "High":
        #             st.error(f"🔴 Risk Level : {level}")
        #
        #         elif level == "Medium":
        #             st.warning(f"🟡 Risk Level : {level}")
        #
        #         else:
        #             st.success(f"🟢 Risk Level : {level}")
        #
        #         # Display Reasons
        #         st.markdown("### 📋 Reasons")
        #
        #         for item in reasons:
        #             st.write(f"✅ {item}")
        #
        #         st.markdown("---")

        # ======================================================
        # Workbook Update + Audit Summary
        # ======================================================

        st.header("📄 Workbook Updater")



        if st.button("Update Workbook",key="worbook_updater_button"):

            if uploaded_file is None:
                st.warning("Please upload a workbook.")
            elif not requirement.strip():
                st.warning("Please enter a business requirement.")
            else:
                with st.spinner("Updating workbook..."):

                    # Generate requirement analysis
                    analysis = analyze_requirement(requirement)

                    # Output file name
                    output_file = "Updated_Production.xlsx"

                    # Call workbook updater
                    update_workbook(
                        uploaded_file,
                        output_file,
                        analysis
                    )

                    st.success("✅ Workbook updated successfully!")

                    with open(output_file, "rb") as f:
                        st.download_button(
                            "📥 Download Updated Workbook",
                            data=f,
                            file_name="Updated_Production.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        )




            # ======================================================
            # Impact Analysis
            # ======================================================


        # ==========================================
        # IMPACT ANALYZER
        # ==========================================

        st.header("🔍 Impact Analyzer")
        run_impact = st.button(
                "Run Impact Analysis",
                key="run_impact_button"
            )

        if run_impact:

                if uploaded_file is None:
                    st.warning("Please upload a workbook.")

                elif not requirement.strip():
                    st.warning("Please enter a business requirement.")

                else:

                    with st.spinner("Running Impact Analysis..."):

                        try:

                            result = impact_analysis(
                                input_file=uploaded_file,
                                requirement=requirement
                            )

                            st.success("✅ Impact Analysis completed successfully.")

                            # st.write("### Summary")
                            # st.write(result.get("summary", ""))
                            #
                            # st.write("### Affected Sheets")
                            # st.write(result.get("affected_sheets", []))
                            #
                            # st.write("### Affected Columns")
                            # st.write(result.get("affected_columns", []))
                            #
                            # st.write("### Affected Formulas")
                            # st.write(result.get("affected_formulas", []))
                            #
                            # st.write("### Risk Level")
                            # st.info(result.get("risk_level", ""))
                            #
                            # st.write("### Recommendation")
                            # st.write(result.get("recommendation", ""))

                            report = f"""
            IMPACT ANALYSIS REPORT

            Business Requirement:
            {requirement}

            Summary:
            {result.get("summary", "")}

            Affected Sheets:
            {', '.join(result.get("affected_sheets", []))}

            Affected Columns:
            {', '.join(result.get("affected_columns", []))}

            Affected Formulas:
            {', '.join(result.get("affected_formulas", []))}

            Risk Level:
            {result.get("risk_level", "")}

            Recommendation:
            {result.get("recommendation", "")}
            """
                            doc = Document()
                            doc.add_heading("Impact Analysis Report", 0)
                            doc.add_paragraph(report)
                            buffer = BytesIO()
                            doc.save(buffer)
                            buffer.seek(0)

                            st.download_button(
                                "📥 Download Impact Analysis Report",
                                data=buffer.getvalue(),
                                file_name="Impact_Analysis_Report.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                key="impact_download"
                            )

                        except Exception as e:
                            st.error(f"Impact Analysis Error: {e}")
        # ==========================================
        # Audit Summary
        # ==========================================

        st.header("📋 Audit Summary")


        if st.button("Generate Audit Summary", key="audit_button"):

            if uploaded_file is None:
                st.warning("Please upload a workbook.")

            elif not requirement.strip():
                st.warning("Please enter a business requirement.")

            else:

                with st.spinner("Generating Audit Summary..."):

                    try:

                        result = generate_audit_summary(
                            input_file=uploaded_file,
                            requirement=requirement
                        )

                        st.success("✅ Audit Summary Generated Successfully")

                        # st.write("### Audit Summary")
                        # st.write(result.get("audit_summary", ""))
                        #
                        # st.write("### Risk Level")
                        # st.info(result.get("risk_level", ""))
                        #
                        # st.write("### Compliance Status")
                        # st.write(result.get("compliance_status", ""))
                        #
                        # st.write("### Recommendation")
                        # st.write(result.get("recommendation", ""))

                        report = f"""
            AUDIT SUMMARY

            Workbook:
            {result.get("workbook_name", "")}

            Business Requirement:
            {requirement}

            Audit Summary:
            {result.get("audit_summary", "")}

            Risk Level:
            {result.get("risk_level", "")}

            Change Type:
            {result.get("change_type", "")}

            Compliance Status:
            {result.get("compliance_status", "")}

            Recommendation:
            {result.get("recommendation", "")}
            """
                        doc = Document()
                        doc.add_heading("Audit Summary Report", 0)
                        doc.add_paragraph(report)
                        buffer = BytesIO()
                        doc.save(buffer)
                        buffer.seek(0)
                        st.download_button(
                            "📥 Download Audit Summary",
                            data=buffer.getvalue(),
                            file_name="Audit_Summary_Report.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key="audit_download"
                        )

                    except Exception as e:
                        st.error(f"Audit Summary Error: {e}")

    # ==========================================
    # CR Generation
    # ==========================================
        st.header("📋 Generate Change Request")
        if st.button("Generate CR", key="cr_button"):

            if uploaded_file is None:
                st.error("Please upload a file first.")

            else:
                impact = "Medium Impact"
                st.write("Result Data:")
                st.json(result)
                cr_file = generate_change_request(
                    result,
                    requirement,
                    uploaded_file.name,
                    impact
                )

                st.success("✅ Change Request document has been generated successfully.")

                with open(cr_file, "rb") as file:
                    # doc = Document()
                    # doc.add_heading("CR Generation", 0)
                    # doc.add_paragraph(result)
                    # buffer = BytesIO()
                    # doc.save(buffer)
                    # buffer.seek(0)

                    st.download_button(
                        label="📄 Download Change Request",
                        data=file,
                        file_name="Change_Request.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )

