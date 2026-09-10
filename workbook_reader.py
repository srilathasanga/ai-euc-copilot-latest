import openpyxl
from openpyxl.cell.cell import Cell
# Read the uploaded workbook
def read_workbook(uploaded_file):
    workbook = openpyxl.load_workbook(uploaded_file)
    return workbook


# Get workbook information
def workbook_information(workbook, uploaded_file):

    info = {
        "Workbook Name": uploaded_file.name,
        "Total Sheets": len(workbook.sheetnames),
        "Sheet Names": workbook.sheetnames,

    }

    return info
# to show sheet names along with existing formulas info in production file
def get_formula_information(workbook):

    formula_info = []

    for sheet_name in workbook.sheetnames:

        ws = workbook[sheet_name]

        for row in ws.iter_rows():
            for cell in row:

                if isinstance(cell.value, str) and cell.value.startswith("="):

                    formula_info.append({
                        "Sheet": sheet_name,
                        "Cell": cell.coordinate,
                        "Formula": cell.value
                    })

    return formula_info