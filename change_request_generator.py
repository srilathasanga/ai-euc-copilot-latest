# from docx import Document
# from datetime import datetime
#
# def generate_change_request(result, business_requirement, file_name, impact):
#
#     doc = Document()
#
#     doc.add_heading("Change Request Document", level=1)
#
#     doc.add_heading("1. General Information", level=2)
#     doc.add_paragraph(f"Date : {datetime.now().strftime('%d-%m-%Y')}")
#     doc.add_paragraph(f"Workbook : {file_name}")
#     doc.add_paragraph(f"Impact : {impact}")
#
#     doc.add_heading("2. Business Requirement", level=2)
#     doc.add_paragraph(business_requirement)
#
#     doc.add_heading("3. Change Details", level=2)
#
#     table = doc.add_table(rows=5, cols=2)
#     table.style = "Table Grid"
#
#     table.cell(0,0).text = "Change Type"
#     table.cell(0,1).text = str(result.get("change_type",""))
#
#     table.cell(1,0).text = "Affected Sheet"
#     table.cell(1,1).text = str(result.get("sheet_name",""))
#
#     table.cell(2,0).text = "Column Position"
#     table.cell(2,1).text = str(result.get("after_column",""))
#
#     table.cell(3,0).text = "New Column"
#     table.cell(3,1).text = str(result.get("column_name",""))
#
#     table.cell(4,0).text = "Formula"
#     table.cell(4,1).text = str(result.get("formula","Not Applicable"))
#
#     output_file = "Change_Request.docx"
#     doc.save(output_file)
#
#     return output_file
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from datetime import datetime


def add_highlighted_text(cell, text):
    """Add yellow highlighted text to a table cell"""
    cell.text = ""

    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def generate_change_request(
        result,
        business_requirement,
        file_name,
        impact):

    doc = Document()

    # ==========================
    # Title
    # ==========================
    doc.add_heading("Change Request Document", level=1)

    # ==========================
    # General Information
    # ==========================
    doc.add_heading("1. General Information", level=2)

    doc.add_paragraph(
        f"Date : {datetime.now().strftime('%d-%m-%Y')}"
    )

    doc.add_paragraph(
        f"Workbook : {file_name}"
    )

    doc.add_paragraph(
        f"Impact : {impact}"
    )

    # ==========================
    # Business Requirement
    # ==========================
    doc.add_heading("2. Business Requirement", level=2)

    doc.add_paragraph(business_requirement)

    # ==========================
    # Change Details
    # ==========================
    doc.add_heading("3. Change Details", level=2)

    operation = result.get("operation", "N/A")
    sheet = result.get("sheet", "N/A")
    column_name = result.get("target", {}).get(
        "column_name",
        "N/A"
    )
    formula = result.get("parameters", {}).get(
        "formula",
        "N/A"
    )

    # Friendly change type
    if operation == "update_formula":
        operation = "Formula Update"
    elif operation == "add_column":
        operation = "Column Addition"
    elif operation == "delete_column":
        operation = "Column Deletion"

    details = [
        ("Change Type", operation),
        ("Affected Sheet", sheet),
        ("Column Name", column_name),
        ("Formula", formula),
        ("Impact", impact)
    ]

    table = doc.add_table(
        rows=len(details),
        cols=2
    )

    table.style = "Table Grid"

    for row_idx, (label, value) in enumerate(details):

        table.cell(row_idx, 0).text = label

        add_highlighted_text(
            table.cell(row_idx, 1),
            value
        )

    # ==========================
    # Save
    # ==========================
    output_file = "Change_Request.docx"

    doc.save(output_file)

    return output_file